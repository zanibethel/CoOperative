# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "pypdf>=5.0.0",
#   "python-docx>=1.1.2",
#   "openpyxl>=3.1.5",
#   "python-pptx>=1.0.2",
# ]
# ///

"""Local document extraction helper for CoOperative Personal AI."""

from __future__ import annotations

import json
import sys
from pathlib import Path

MAX_CHARS = 240_000
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".jsonl", ".xml", ".html", ".htm",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".css", ".scss", ".sql", ".ps1", ".sh",
    ".bat", ".cmd", ".java", ".cs", ".cpp", ".c", ".h", ".hpp", ".go", ".rs",
    ".rb", ".php", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".env", ".log",
}


def trim(value: str) -> str:
    value = value.replace("\x00", "")
    return value[:MAX_CHARS]


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return trim(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return trim(data.decode("utf-8", errors="replace"))


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for index, page in enumerate(reader.pages[:250], start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"--- Page {index} ---\n{text}")
        if sum(len(part) for part in parts) >= MAX_CHARS:
            break
    return trim("\n\n".join(parts))


def extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return trim("\n".join(parts))


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets[:30]:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                parts.append("\t".join(values))
            if sum(len(part) for part in parts) >= MAX_CHARS:
                return trim("\n".join(parts))
    return trim("\n".join(parts))


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides[:200], start=1):
        slide_parts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str) and text.strip():
                slide_parts.append(text)
        if slide_parts:
            parts.append(f"--- Slide {index} ---\n" + "\n".join(slide_parts))
        if sum(len(part) for part in parts) >= MAX_CHARS:
            break
    return trim("\n\n".join(parts))


def extract(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = extract_pdf(path)
        kind = "pdf"
    elif suffix == ".docx":
        text = extract_docx(path)
        kind = "document"
    elif suffix in {".xlsx", ".xlsm"}:
        text = extract_xlsx(path)
        kind = "spreadsheet"
    elif suffix == ".pptx":
        text = extract_pptx(path)
        kind = "presentation"
    elif suffix in TEXT_EXTENSIONS or not suffix:
        text = read_text(path)
        kind = "text"
    else:
        raise RuntimeError(
            "Unsupported file type. Supported: PDF, DOCX, XLSX/XLSM, PPTX, text, "
            "Markdown, CSV/JSON, and common code/config files."
        )

    if not text.strip():
        raise RuntimeError("No readable text was found in this file.")

    return {"kind": kind, "text": text, "characters": len(text)}


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: windows-local-file-extract.py <file>")
    path = Path(sys.argv[1]).resolve()
    if not path.is_file():
        raise RuntimeError("File does not exist.")
    print(json.dumps(extract(path), ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"error": str(exc)[:800]}))
        raise SystemExit(1)
